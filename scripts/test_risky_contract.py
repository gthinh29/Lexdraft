"""
scripts/test_risky_contract.py

Kiểm thử định lượng khả năng phát hiện rủi ro và điều khoản vi phạm
trên bộ 5 hợp đồng đối sánh (2 Hợp đồng chuẩn sạch + 3 Hợp đồng cài cắm rủi ro).

Danh sách 5 hợp đồng:
1. mau_hop_dong_thiet_ke_website.md (Clean 1, 15 điều, mẫu chuẩn đã tinh chỉnh Điều 11 & 12, 0 lỗi kỳ vọng)
2. test_contract_consulting_clean.docx (Clean 2, 8 điều, HĐ dịch vụ tư vấn ERP chuẩn pháp lý, 0 lỗi kỳ vọng)
3. test_contract_risky.docx (Risky 1, 7 điều, HĐ phần mềm với 5 lỗi cài cắm)
4. test_contract_risky_2.docx (Risky 2, 7 điều, HĐ bảo trì máy chủ với 4 lỗi cài cắm)
5. test_contract_risky_3.docx (Risky 3, 6 điều, HĐ Digital Marketing & SEO với 4 lỗi cài cắm)

Tổng số lỗi cài cắm (Ground Truth Faults): 13 lỗi
Tính toán:
- Tỉ lệ phát hiện lỗi (Recall / Sensitivity) = TP / (TP + FN)
- Tỉ lệ cảnh báo sai (False Positive Rate) trên các hợp đồng sạch / điều khoản sạch
- Báo cáo chi tiết từng hợp đồng ra data/eval/contract_benchmark_report.json
"""

import json
import logging
import sys
import time
from pathlib import Path

# Thêm thư mục gốc vào path
PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

# Cấu hình UTF-8 cho Windows console
if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("contract_benchmark")

import docx  # noqa: E402
from modules.risk_assessment.service import analyze_contract  # noqa: E402


def generate_benchmark_contracts():
    """Tự động tạo các file docx hợp đồng mẫu nếu chưa tồn tại."""
    data_dir = PROJECT_ROOT / "data"
    data_dir.mkdir(parents=True, exist_ok=True)

    # 1. Clean 2: test_contract_consulting_clean.docx (8 điều khoản chuẩn chỉ)
    clean_2_path = data_dir / "test_contract_consulting_clean.docx"
    if not clean_2_path.exists():
        doc = docx.Document()
        doc.add_heading(
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n---o0o---",
            level=2,
        )
        doc.add_heading(
            "HỢP ĐỒNG DỊCH VỤ TƯ VẤN QUẢN TRỊ DOANH NGHIỆP VÀ TRIỂN KHAI ERP", level=1
        )
        doc.add_paragraph("Số: 26/2026/HĐDV-ERP\n")
        doc.add_paragraph(
            "Căn cứ Bộ luật Dân sự số 91/2015/QH13;\nCăn cứ Luật Thương mại số 36/2005/QH11;\nCăn cứ Luật Sở hữu trí tuệ số 50/2005/QH11 (sửa đổi, bổ sung năm 2022);"
        )

        doc.add_paragraph(
            "Hôm nay, ngày 20 tháng 02 năm 2026, tại Hà Nội, chúng tôi gồm có:\n"
            "BÊN THUÊ DỊCH VỤ (BÊN A): CÔNG TY CỔ PHẦN DƯỢC PHẨM AN BÌNH\n"
            "Đại diện: Ông Nguyễn Minh Long - Chức vụ: Tổng Giám đốc\n"
            "BÊN CUNG CẤP DỊCH VỤ (BÊN B): CÔNG TY TNHH GIẢI PHÁP & TƯ VẤN MINH TRIẾT\n"
            "Đại diện: Bà Lê Thanh Hằng - Chức vụ: Giám đốc"
        )

        articles = [
            (
                "Điều 1. Phạm vi công việc và nội dung dịch vụ",
                "Bên B nhận cung cấp dịch vụ tư vấn quy trình quản trị tài chính, chuỗi cung ứng và triển khai phần mềm hoạch định nguồn lực doanh nghiệp (ERP) cho Bên A theo phụ lục mô tả kỹ thuật kèm theo.",
            ),
            (
                "Điều 2. Phí dịch vụ và phương thức thanh toán",
                "Tổng giá trị hợp đồng dịch vụ là 350.000.000 VNĐ (Ba trăm năm mươi triệu đồng), đã bao gồm thuế GTGT. Bên A thanh toán thành 03 đợt tương ứng theo tiến độ nghiệm thu từng giai đoạn bằng hình thức chuyển khoản.",
            ),
            (
                "Điều 3. Quyền và nghĩa vụ của Bên A",
                "1. Cung cấp đầy đủ, kịp thời tài liệu, số liệu và hạ tầng cần thiết cho Bên B.\n2. Phối hợp cử nhân sự tham gia các buổi khảo sát và đào tạo sử dụng hệ thống.\n3. Thanh toán phí dịch vụ đúng hạn theo thỏa thuận tại Điều 2.",
            ),
            (
                "Điều 4. Quyền và nghĩa vụ của Bên B",
                "1. Bố trí đội ngũ chuyên gia tư vấn có đủ trình độ chuyên môn để thực hiện dịch vụ đúng cam kết.\n2. Bàn giao báo cáo phân tích và tài liệu hướng dẫn vận hành đúng tiến độ.\n3. Giữ bảo mật tuyệt đối toàn bộ số liệu tài chính và hoạt động kinh doanh của Bên A theo thỏa thuận bảo mật.",
            ),
            (
                "Điều 5. Phạt vi phạm và bồi thường thiệt hại",
                "1. Trong trường hợp Bên B chậm trễ bàn giao dịch vụ quá 10 ngày làm việc mà không có lý do chính đáng, Bên B phải chịu phạt vi phạm 0,1% giá trị phần nghĩa vụ bị vi phạm cho mỗi ngày chậm trễ, nhưng tổng mức phạt không vượt quá 8% giá trị phần nghĩa vụ hợp đồng bị vi phạm theo đúng quy định tại Điều 301 Luật Thương mại 2005.\n2. Bên vi phạm nghĩa vụ hợp đồng có trách nhiệm bồi thường toàn bộ thiệt hại thực tế phát sinh trực tiếp theo Điều 302 Luật Thương mại 2005.",
            ),
            (
                "Điều 6. Sự kiện bất khả kháng",
                "Trường hợp một bên không thể thực hiện nghĩa vụ do sự kiện bất khả kháng (thiên tai, dịch bệnh, quyết định cấm của cơ quan nhà nước có thẩm quyền), bên bị ảnh hưởng phải thông báo bằng văn bản trong vòng 07 ngày và được miễn trừ trách nhiệm theo quy định tại Điều 294 Luật Thương mại 2005 và Điều 156 Bộ luật Dân sự 2015.",
            ),
            (
                "Điều 7. Quyền sở hữu trí tuệ và bảo mật",
                "1. Toàn bộ tài liệu quy trình được xây dựng tùy chỉnh riêng cho Bên A thuộc quyền sở hữu của Bên A.\n2. Bên B giữ quyền tác giả và quyền sở hữu đối với các module phần mềm nền tảng gốc, đồng thời cấp quyền sử dụng không độc quyền, vô thời hạn cho Bên A theo quy định của Luật Sở hữu trí tuệ.",
            ),
            (
                "Điều 8. Giải quyết tranh chấp và hiệu lực thi hành",
                "Mọi tranh chấp phát sinh sẽ được giải quyết trước hết bằng thương lượng, hòa giải. Nếu không đạt được thỏa thuận, vụ việc sẽ được đưa ra giải quyết tại Trung tâm Trọng tài Quốc tế Việt Nam (VIAC) theo Quy tắc tố tụng trọng tài của VIAC. Hợp đồng có hiệu lực kể từ ngày ký.",
            ),
        ]
        for title, content in articles:
            doc.add_heading(title, level=2)
            doc.add_paragraph(content)
        doc.save(clean_2_path)
        logger.info(f"Đã tạo Hợp đồng Clean 2: {clean_2_path}")

    # 2. Risky 2: test_contract_risky_2.docx (Bảo trì Cloud, 4 lỗi cài cắm)
    risky_2_path = data_dir / "test_contract_risky_2.docx"
    if not risky_2_path.exists():
        doc = docx.Document()
        doc.add_heading(
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n---o0o---",
            level=2,
        )
        doc.add_heading(
            "HỢP ĐỒNG DỊCH VỤ QUẢN TRỊ VÀ BẢO TRÌ HỆ THỐNG MÁY CHỦ CLOUD", level=1
        )
        doc.add_paragraph("Số: 88/2026/HĐDV-CLOUD\n")
        # LỖI 1: Căn cứ Luật Thương mại 2020 (Luật giả mạo / không tồn tại)
        doc.add_paragraph(
            "Căn cứ Bộ luật Dân sự số 91/2015/QH13;\nCăn cứ Luật Thương mại 2020;\nCăn cứ Luật An toàn thông tin mạng số 86/2015/QH13;"
        )

        doc.add_paragraph(
            "Hôm nay, ngày 10 tháng 03 năm 2026, tại TP. Hồ Chí Minh:\n"
            "BÊN THUÊ (BÊN A): CÔNG TY CỔ PHẦN FINTECH VIỆT\n"
            "BÊN CUNG CẤP (BÊN B): CÔNG TY TNHH HẠ TẦNG CLOUD SAO MAI"
        )

        articles_risky_2 = [
            (
                "Điều 1. Phạm vi công việc",
                "Bên B cung cấp dịch vụ quản trị, giám sát hệ thống cụm máy chủ ảo hóa đám mây (Cloud Server) và sao lưu định kỳ cho Bên A.",
            ),
            (
                "Điều 2. Chi phí dịch vụ và thanh toán",
                "Phí dịch vụ cố định hàng tháng là 30.000.000 VNĐ. Bên A thanh toán định kỳ vào ngày 05 đầu mỗi tháng.",
            ),
            (
                "Điều 3. Quyền đơn phương chấm dứt hợp đồng",
                "Bên A có toàn quyền đơn phương chấm dứt hợp đồng này tại bất kỳ thời điểm nào mà không cần phải thông báo trước cho Bên B và không phải thanh toán bất kỳ khoản bồi thường hay chi phí phát sinh nào.",
            ),  # LỖI 2: Vi phạm Điều 520 BLDS 2015 (chấm dứt không báo trước thời hạn hợp lý)
            (
                "Điều 4. Chế tài phạt vi phạm hợp đồng",
                "Nếu Bên B để xảy ra sự cố ngừng trệ hệ thống vượt quá 02 giờ trong một tháng, Bên B phải chịu mức phạt vi phạm là 20% tổng giá trị hợp đồng năm.",
            ),  # LỖI 3: Phạt 20% vượt quá trần 8% theo Điều 301 LTM 2005
            (
                "Điều 5. Quản lý và xử lý dữ liệu khách hàng",
                "Bên B được quyền tự do sao lưu, phân tích và chuyển giao toàn bộ cơ sở dữ liệu khách hàng, nhật ký giao dịch của Bên A cho các đối tác công nghệ thứ ba tại nước ngoài phục vụ việc huấn luyện mô hình AI mà không cần phải có văn bản chấp thuận riêng của Bên A hoặc của các chủ thể dữ liệu cá nhân.",
            ),  # LỖI 4: Vi phạm nghiêm trọng Nghị định 13/2023/NĐ-CP và Luật Bảo vệ dữ liệu cá nhân 2025 về chuyển dữ liệu ra nước ngoài và sự đồng ý
            (
                "Điều 6. Cam kết bảo mật thông tin",
                "Hai bên cam kết giữ bí mật các thông tin nội bộ của nhau trong suốt quá trình thực hiện hợp đồng.",
            ),
            (
                "Điều 7. Điều khoản thi hành",
                "Hợp đồng có hiệu lực kể từ ngày ký. Mọi tranh chấp được giải quyết tại Tòa án nhân dân có thẩm quyền.",
            ),
        ]
        for title, content in articles_risky_2:
            doc.add_heading(title, level=2)
            doc.add_paragraph(content)
        doc.save(risky_2_path)
        logger.info(f"Đã tạo Hợp đồng Risky 2: {risky_2_path}")

    # 3. Risky 3: test_contract_risky_3.docx (Digital Marketing & SEO, 4 lỗi cài cắm)
    risky_3_path = data_dir / "test_contract_risky_3.docx"
    if not risky_3_path.exists():
        doc = docx.Document()
        doc.add_heading(
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM\nĐộc lập - Tự do - Hạnh phúc\n---o0o---",
            level=2,
        )
        doc.add_heading(
            "HỢP ĐỒNG DỊCH VỤ DIGITAL MARKETING VÀ TỐI ƯU HÓA CÔNG CỤ TÌM KIẾM (SEO)",
            level=1,
        )
        doc.add_paragraph("Số: 99/2026/HĐDV-MKT\n")
        # LỖI 1: Căn cứ Luật Bảo vệ quyền lợi người tiêu dùng số 59/2010/QH12 (hết hiệu lực từ 01/07/2024)
        doc.add_paragraph(
            "Căn cứ Bộ luật Dân sự số 91/2015/QH13;\nCăn cứ Luật Thương mại số 36/2005/QH11;\nCăn cứ Luật Bảo vệ quyền lợi người tiêu dùng số 59/2010/QH12;"
        )

        doc.add_paragraph(
            "Hôm nay, ngày 05 tháng 04 năm 2026, tại TP. Hà Nội:\n"
            "BÊN THUÊ (BÊN A): CÔNG TY CỔ PHẦN THỜI TRANG SMARTLOOK\n"
            "BÊN CUNG CẤP (BÊN B): CÔNG TY TNHH TRUYỀN THÔNG & QUẢNG CÁO LEADMAX"
        )

        articles_risky_3 = [
            (
                "Điều 1. Nội dung dịch vụ và phạm vi triển khai",
                "Bên B chịu trách nhiệm triển khai chiến dịch quảng cáo kỹ thuật số, tối ưu hóa công cụ tìm kiếm (SEO) đưa 20 từ khóa sản phẩm lên trang nhất Google và xây dựng tệp dữ liệu khách hàng tiềm năng cho Bên A.",
            ),
            (
                "Điều 2. Thu thập và khai thác dữ liệu người dùng",
                "Để phục vụ chiến dịch tiếp thị và bán hàng trực tiếp (telesale), Bên B sẽ sử dụng các công cụ tự động quét, trích xuất số điện thoại, email và vị trí địa lý của mọi người dùng truy cập website của Bên A mà không cần hiển thị thông báo hay lấy sự đồng ý của người dùng.",
            ),  # LỖI 2: Vi phạm nghiêm trọng Luật Bảo vệ dữ liệu cá nhân 2025 và Luật GDĐT 2023
            (
                "Điều 3. Cam kết KPI và chế tài phạt vi phạm",
                "Bên B cam kết đạt tối thiểu 1.000 đơn hàng chuyển đổi/tháng. Nếu kết thúc quý tỉ lệ đạt dưới 70% chỉ tiêu KPI đã cam kết, Bên B phải chịu phạt vi phạm bằng 25% tổng phí dịch vụ của toàn quý đó.",
            ),  # LỖI 3: Mức phạt 25% vi phạm trần 8% theo Điều 301 LTM 2005
            (
                "Điều 4. Miễn trừ trách nhiệm pháp lý về quảng cáo",
                "Bên B hoàn toàn được miễn trừ mọi trách nhiệm hành chính, dân sự hoặc hình sự nếu hình ảnh, video quảng cáo và thông điệp truyền thông do Bên B thiết kế bị cơ quan quản lý nhà nước xử phạt vi phạm thuần phong mỹ tục, vi phạm Luật Quảng cáo hoặc xâm phạm bản quyền sở hữu trí tuệ của bất kỳ bên thứ ba nào.",
            ),  # LỖI 4: Miễn trừ trách nhiệm tuyệt đối trái luật theo Điều 351, 360 BLDS 2015, Luật Quảng cáo, Luật SHTT
            (
                "Điều 5. Chi phí dịch vụ và thanh toán",
                "Tổng kinh phí dịch vụ là 80.000.000 VNĐ/tháng. Bên A thanh toán thành từng đợt vào đầu mỗi tháng.",
            ),
            (
                "Điều 6. Điều khoản giải quyết tranh chấp",
                "Hợp đồng có hiệu lực 01 năm kể từ ngày ký. Mọi tranh chấp phát sinh được đưa ra Tòa án nhân dân TP. Hà Nội giải quyết.",
            ),
        ]
        for title, content in articles_risky_3:
            doc.add_heading(title, level=2)
            doc.add_paragraph(content)
        doc.save(risky_3_path)
        logger.info(f"Đã tạo Hợp đồng Risky 3: {risky_3_path}")


def run_benchmark():
    """Chạy đánh giá phân tích rủi ro trên toàn bộ 5 hợp đồng đối sánh."""
    generate_benchmark_contracts()

    benchmark_suite = [
        {
            "id": "CLEAN_1",
            "name": "HĐ Thiết kế Website (Mẫu chuẩn Thư Viện Pháp Luật)",
            "file_path": "data/raw/templates/.md/mau_hop_dong_thiet_ke_website.md",
            "type": "clean",
            "expected_faults": [],
            "total_articles": 15,
        },
        {
            "id": "CLEAN_2",
            "name": "HĐ Dịch vụ Tư vấn ERP (Mẫu chuẩn)",
            "file_path": "data/test_contract_consulting_clean.docx",
            "type": "clean",
            "expected_faults": [],
            "total_articles": 8,
        },
        {
            "id": "RISKY_1",
            "name": "HĐ Phát triển Phần mềm (Risky 1)",
            "file_path": "data/test_contract_risky.docx",
            "type": "risky",
            "expected_faults": [
                {"category": "expired_law", "detail": "Luật CNTT 2006 hết hiệu lực"},
                {"category": "expired_law", "detail": "Luật GDĐT 2005 hết hiệu lực"},
                {
                    "category": "clause_risk",
                    "article": "Điều 4",
                    "detail": "Phạt vi phạm 15% > 8%",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 5",
                    "detail": "Miễn trừ trách nhiệm bồi thường tuyệt đối",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 6",
                    "detail": "Bên B chiếm trọn bản quyền mã nguồn",
                },
            ],
            "total_articles": 7,
        },
        {
            "id": "RISKY_2",
            "name": "HĐ Quản trị & Bảo trì Máy chủ Cloud (Risky 2)",
            "file_path": "data/test_contract_risky_2.docx",
            "type": "risky",
            "expected_faults": [
                {
                    "category": "expired_law",
                    "detail": "Luật Thương mại 2020 giả mạo / không tồn tại",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 3",
                    "detail": "Đơn phương chấm dứt không báo trước",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 4",
                    "detail": "Phạt vi phạm 20% > 8%",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 5",
                    "detail": "Chuyển dữ liệu cá nhân ra nước ngoài trái luật",
                },
            ],
            "total_articles": 7,
        },
        {
            "id": "RISKY_3",
            "name": "HĐ Digital Marketing & SEO (Risky 3)",
            "file_path": "data/test_contract_risky_3.docx",
            "type": "risky",
            "expected_faults": [
                {"category": "expired_law", "detail": "Luật BVQLNTD 2010 hết hiệu lực"},
                {
                    "category": "clause_risk",
                    "article": "Điều 2",
                    "detail": "Quét thu thập dữ liệu người dùng trái phép",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 3",
                    "detail": "Phạt KPI 25% > 8%",
                },
                {
                    "category": "clause_risk",
                    "article": "Điều 4",
                    "detail": "Miễn trừ trách nhiệm vi phạm quảng cáo trái luật",
                },
            ],
            "total_articles": 6,
        },
    ]

    total_injected_faults = 13
    total_tp = 0
    total_fn = 0
    total_clean_articles = 23  # 15 + 8
    total_fp = 0

    benchmark_results = []

    logger.info("==================================================================")
    logger.info("BẮT ĐẦU CHẠY BENCHMARK ĐỐI SÁNH 5 HỢP ĐỒNG (2 CLEAN + 3 RISKY)")
    logger.info("==================================================================")

    for item in benchmark_suite:
        contract_path = PROJECT_ROOT / item["file_path"]
        logger.info(f"\n---> Đang phân tích: {item['name']} ({contract_path.name})")
        start_t = time.time()

        res = analyze_contract(contract_path)
        elapsed = round(time.time() - start_t, 2)

        expired_alerts = res.get("expired_law_alerts", [])
        risk_results = res.get("risk_results", [])

        # Đánh giá cảnh báo
        detected_fault_summaries = []
        for alert in expired_alerts:
            detected_fault_summaries.append(
                {
                    "type": "expired_law",
                    "law": alert.get("law_name", ""),
                    "status": alert.get("status", ""),
                    "replacement": alert.get("replacement_law", ""),
                    "note": alert.get("note", ""),
                }
            )

        critical_risks = []
        for r in risk_results:
            rui_ro_text = r.get("rui_ro", "")
            dieu_khoan_title = r.get("dieu_khoan", "")
            # Phân loại rủi ro nếu có phát hiện rủi ro thực chất (loại trừ phản hồi an toàn)
            is_risk = False
            r_lower = rui_ro_text.lower()
            if any(
                k in r_lower
                for k in [
                    "rủi ro",
                    "vi phạm",
                    "không phù hợp",
                    "vượt quá",
                    "trái quy định",
                    "cần sửa",
                    "vô hiệu",
                ]
            ):
                if not any(
                    safe_k in r_lower
                    for safe_k in [
                        "không phát hiện rủi ro",
                        "phù hợp quy định",
                        "tuân thủ đầy đủ",
                        "không có rủi ro đáng kể",
                    ]
                ):
                    is_risk = True

            if is_risk:
                critical_risks.append(
                    {
                        "dieu_khoan": dieu_khoan_title[:50],
                        "rui_ro": rui_ro_text[:150] + "...",
                        "can_cu": r.get("can_cu", []),
                    }
                )

        item_res = {
            "id": item["id"],
            "name": item["name"],
            "type": item["type"],
            "total_articles": item["total_articles"],
            "elapsed_seconds": elapsed,
            "detected_expired_alerts": len(expired_alerts),
            "detected_critical_risks": len(critical_risks),
            "expired_details": detected_fault_summaries,
            "risk_details": critical_risks,
        }

        if item["type"] == "clean":
            # Trên hợp đồng chuẩn: đếm số cảnh báo sai (False Positive)
            fp_count = len(expired_alerts) + len(critical_risks)
            total_fp += fp_count
            item_res["false_positives"] = fp_count
            logger.info(
                f"Kết quả {item['name']}: {fp_count} cảnh báo sai (FP) / {item['total_articles']} điều khoản."
            )
        else:
            # Trên hợp đồng risky: đếm True Positives so với expected_faults
            injected_count = len(item["expected_faults"])
            tp_found = 0
            for ef in item["expected_faults"]:
                found = False
                if ef["category"] == "expired_law":
                    for a in expired_alerts:
                        if any(
                            term in a.get("law_name", "").lower()
                            or term in a.get("note", "").lower()
                            for term in ef["detail"].lower().split()
                        ):
                            found = True
                            break
                    # Hoặc có thể được bắt trong risk_results
                    if not found:
                        for cr in critical_risks:
                            if any(
                                term in cr["rui_ro"].lower()
                                for term in ef["detail"].lower().split()
                            ):
                                found = True
                                break
                else:
                    for cr in critical_risks:
                        art = ef.get("article", "").lower()
                        if art in cr["dieu_khoan"].lower() or any(
                            term in cr["rui_ro"].lower()
                            for term in ef["detail"].lower().split()
                            if len(term) > 3
                        ):
                            found = True
                            break
                if found:
                    tp_found += 1

            total_tp += tp_found
            fn = injected_count - tp_found
            total_fn += fn
            item_res["injected_faults"] = injected_count
            item_res["true_positives"] = tp_found
            item_res["false_negatives"] = fn
            logger.info(
                f"Kết quả {item['name']}: Phát hiện {tp_found}/{injected_count} lỗi cài cắm (Bỏ sót: {fn})."
            )

        benchmark_results.append(item_res)

    # Tính toán chỉ số tổng hợp
    recall = (
        round(total_tp / total_injected_faults, 4) if total_injected_faults > 0 else 0.0
    )
    fp_rate = (
        round(total_fp / total_clean_articles, 4) if total_clean_articles > 0 else 0.0
    )
    specificity = round(1.0 - fp_rate, 4)

    summary = {
        "benchmark_suite_size": len(benchmark_suite),
        "total_injected_faults": total_injected_faults,
        "total_true_positives": total_tp,
        "total_false_negatives": total_fn,
        "fault_detection_recall": recall,
        "total_clean_articles_tested": total_clean_articles,
        "total_false_positives": total_fp,
        "false_positive_rate": fp_rate,
        "specificity": specificity,
        "contract_results": benchmark_results,
    }

    # Xuất báo cáo
    report_file = PROJECT_ROOT / "data" / "eval" / "contract_benchmark_report.json"
    report_file.parent.mkdir(parents=True, exist_ok=True)
    with open(report_file, "w", encoding="utf-8") as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)

    logger.info("\n==================================================================")
    logger.info("BÁO CÁO TỔNG KẾT BENCHMARK 5 HỢP ĐỒNG:")
    logger.info("==================================================================")
    logger.info(f"- Tổng số lỗi cài cắm: {total_injected_faults}")
    logger.info(f"- Số lỗi phát hiện chính xác (TP): {total_tp}")
    logger.info(f"- Số lỗi bỏ sót (FN): {total_fn}")
    logger.info(f"- Tỉ lệ bao phủ phát hiện lỗi (Recall): {recall * 100:.2f}%")
    logger.info(f"- Tổng số điều khoản sạch kiểm thử: {total_clean_articles}")
    logger.info(f"- Số cảnh báo sai (FP): {total_fp}")
    logger.info(
        f"- Độ đặc hiệu trên điều khoản chuẩn (Specificity): {specificity * 100:.2f}%"
    )
    logger.info(f"- Đã lưu báo cáo chi tiết ra: {report_file}")
    logger.info("==================================================================")

    return summary


if __name__ == "__main__":
    run_benchmark()
